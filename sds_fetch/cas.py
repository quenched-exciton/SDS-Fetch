"""
cas.py - Reading and checking CAS Registry Numbers.

A CAS Registry Number has three parts separated by hyphens:

    7732-18-5
    ^^^^ ^^ ^
    |    |  +-- 1 check digit
    |    +----- 2 digits
    +---------- 2 to 7 digits

The check digit lets us catch most typing mistakes without going online.
Take every digit except the check digit, read them from RIGHT to LEFT,
multiply the 1st by 1, the 2nd by 2, the 3rd by 3 and so on, add the
products, and keep only the last digit of the sum (the sum modulo 10).
That last digit must equal the check digit.

Worked example for water, 7732-18-5:
    digits right-to-left : 8  1  2  3  7  7
    multipliers          : 1  2  3  4  5  6
    products             : 8  2  6 12 35 42  -> sum = 105 -> 105 % 10 = 5  (valid)

The list the user types has one chemical per line:

    7758-99-8                        <- just a CAS number
    2530-83-8 @ Gelest               <- "@" + manufacturer whose SDS is wanted
    81-07-2   Saccharin              <- other text on the line is ignored
    # Nickel bath additives          <- "#" starts a note; the rest of the line is ignored

This module has no network code, so everything in it can be tested offline.
"""

# "from __future__ import annotations" lets us write modern type hints such as
# "list[str]" while still supporting older Python versions (3.9).
from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

# ---------------------------------------------------------------------------
# Regular expressions (text patterns) used below
# ---------------------------------------------------------------------------

# Word processors and PDF files often replace the normal hyphen "-" with look-alike
# dash characters (en dash, em dash, minus sign, ...). We convert all of them to a
# plain hyphen before searching for CAS numbers.
_DASH_CHARACTERS = "‐‑‒–—―−﹣－"
_DASH_TRANSLATION = str.maketrans({dash: "-" for dash in _DASH_CHARACTERS})

# Pattern for "something that looks like a CAS number":
#   (\d{2,7})  -> 2 to 7 digits   (first part, captured as group 1)
#   -(\d{2})   -> hyphen + 2 digits (second part, group 2)
#   -(\d)      -> hyphen + 1 digit  (check digit, group 3)
# The "(?<![\d-])" and "(?![\d-])" parts make sure the match is not glued to more
# digits or hyphens on either side (so "164-17-55" is NOT read as "64-17-5").
_CAS_PATTERN = re.compile(r"(?<![\d-])(\d{2,7})-(\d{2})-(\d)(?![\d-])")

# Header words that mark the manufacturer column in a CSV or Excel file.
_MANUFACTURER_HEADERS = ("manufacturer", "supplier", "brand", "vendor", "maker")

# Text that replaces an Excel cell in which Excel turned a CAS number into a date.
EXCEL_DATE_MARK = "Excel date"


# ---------------------------------------------------------------------------
# Small data containers
# ---------------------------------------------------------------------------

@dataclass
class InvalidEntry:
    """One piece of user input that could not be used as a CAS number."""

    text: str    # exactly what the user typed (or what was in the CSV cell)
    reason: str  # human-readable explanation, written to the log file
    line: int = 0  # line number in the list (1 = first line); 0 = unknown


@dataclass
class ParsedInput:
    """Everything we learned from reading the user's list."""

    valid_cas: list[str] = field(default_factory=list)          # unique, in input order
    invalid: list[InvalidEntry] = field(default_factory=list)   # rejected entries
    duplicates_ignored: int = 0                                  # repeats that were skipped
    entries_read: int = 0                                        # non-empty lines / cells seen
    # Manufacturer requested for a CAS number with "@", e.g. {"2530-83-8": "Gelest"}.
    manufacturers: dict[str, str] = field(default_factory=dict)
    duplicate_lines: list[int] = field(default_factory=list)     # line numbers of repeats


# ---------------------------------------------------------------------------
# Core CAS helpers
# ---------------------------------------------------------------------------

def normalize_dashes(text: str) -> str:
    """Replace every look-alike dash character with a plain hyphen '-'."""
    return text.translate(_DASH_TRANSLATION)


def has_valid_check_digit(cas: str) -> bool:
    """
    Return True if the CAS number's check digit is correct.

    `cas` must already be in the form "digits-digits-digit" (e.g. "64-17-5").
    """
    first_part, second_part, check_digit = cas.split("-")

    # Join the first two parts ("64" + "17" -> "6417") and reverse the string
    # ("7146") so that position 1 is the right-most digit.
    digits_right_to_left = (first_part + second_part)[::-1]

    # enumerate(..., start=1) gives (1, '7'), (2, '1'), (3, '4'), (4, '6')
    total = sum(position * int(digit)
                for position, digit in enumerate(digits_right_to_left, start=1))

    return total % 10 == int(check_digit)


def check_cas(text: str) -> tuple[str | None, str]:
    """
    Check a single piece of text that should be exactly one CAS number.

    Returns a pair (cas, message):
      * ("64-17-5", "")               when the text is a valid CAS number
      * (None, "reason it is invalid") otherwise
    """
    cleaned = normalize_dashes(text.strip())
    match = _CAS_PATTERN.fullmatch(cleaned)
    if match is None:
        return None, "not in CAS format (expected e.g. 64-17-5)"
    return _check_match(match)


def _check_match(match: re.Match) -> tuple[str | None, str]:
    """Validate one regular-expression match of _CAS_PATTERN (see check_cas)."""
    first_part, second_part, check_digit = match.groups()

    # Remove leading zeros from the first part: "064-17-5" -> "64-17-5".
    # CAS numbers never start with zero, but spreadsheets sometimes add them.
    first_part = first_part.lstrip("0")
    if len(first_part) < 2:
        return None, "first part of a CAS number must have 2 to 7 digits"

    cas = f"{first_part}-{second_part}-{check_digit}"
    if not has_valid_check_digit(cas):
        return None, "check digit is wrong (probably a typo)"
    return cas, ""


def extract_cas_numbers(text: str) -> tuple[list[str], list[InvalidEntry]]:
    """
    Find every CAS-looking number inside one line of text (or one CSV cell).

    This is forgiving on purpose: a pasted line such as
        "64-17-5   Ethanol, absolute"
    gives ["64-17-5"], and the chemical name is simply ignored.

    Returns (valid_cas_list, invalid_entries).
    If the text contains nothing that looks like a CAS number, the whole text is
    reported as one invalid entry so the user can see it in the log.
    """
    cleaned = normalize_dashes(text.strip())
    found_valid: list[str] = []
    found_invalid: list[InvalidEntry] = []

    if cleaned.startswith(EXCEL_DATE_MARK):
        found_invalid.append(InvalidEntry(cleaned, "Excel turned this CAS number into a date: retype it "
                                                   "(in Excel, format the CAS column as Text)"))
        return found_valid, found_invalid

    matches = list(_CAS_PATTERN.finditer(cleaned))
    if not matches:
        found_invalid.append(InvalidEntry(text.strip(), "no CAS number found (expected e.g. 64-17-5)"))
        return found_valid, found_invalid

    for match in matches:
        cas, reason = _check_match(match)
        if cas is None:
            found_invalid.append(InvalidEntry(match.group(0), reason))
        else:
            found_valid.append(cas)
    return found_valid, found_invalid


def split_line(line: str) -> tuple[str, str]:
    """
    Split one line of the list into (CAS part, manufacturer).

        "2530-83-8 @ Gelest   # epoxy silane"  ->  ("2530-83-8 ", "Gelest")
        "7758-99-8"                            ->  ("7758-99-8", "")

    The note after "#" is dropped. The manufacturer ends at a tab, so a line
    pasted from three Excel columns still gives a clean name.
    """
    line = line.split("#", 1)[0]
    if "@" not in line:
        return line, ""
    cas_part, manufacturer = line.split("@", 1)
    manufacturer = manufacturer.split("\t", 1)[0].strip(" ,;")
    return cas_part, manufacturer


# ---------------------------------------------------------------------------
# Reading the user's list
# ---------------------------------------------------------------------------

def parse_text_list(text: str) -> ParsedInput:
    """
    Read the list typed or pasted into the window (format: see top of this file).

    Several CAS numbers on one line also work when they are separated by commas,
    semicolons, tabs or spaces; an "@ manufacturer" then applies to all of them.
    Duplicate CAS numbers are kept only once, in the order they first appeared.
    """
    result = ParsedInput()
    already_seen: set[str] = set()

    # enumerate(..., start=1) numbers the lines the way an editor does: 1, 2, 3, ...
    for line_number, line in enumerate(text.splitlines(), start=1):
        cas_part, manufacturer = split_line(line)
        if not cas_part.strip():
            if manufacturer:  # "@ Gelest" with the CAS number forgotten
                result.entries_read += 1
                result.invalid.append(InvalidEntry(line.split("#", 1)[0].strip(),
                                                   "no CAS number before the @", line_number))
            continue  # otherwise: blank line, or a line that is only a "# note"
        result.entries_read += 1

        valid, invalid = extract_cas_numbers(cas_part)
        for entry in invalid:
            entry.line = line_number
        result.invalid.extend(invalid)
        for cas in valid:
            if cas in already_seen:
                result.duplicates_ignored += 1
                result.duplicate_lines.append(line_number)
                # A later line may add a manufacturer the first one did not give.
                if manufacturer and cas not in result.manufacturers:
                    result.manufacturers[cas] = manufacturer
                continue
            already_seen.add(cas)
            result.valid_cas.append(cas)
            if manufacturer:
                result.manufacturers[cas] = manufacturer
    return result


# ---------------------------------------------------------------------------
# Loading a CSV or Excel file into the list
# ---------------------------------------------------------------------------

def _read_csv_text(csv_path: Path) -> str:
    """
    Read a CSV file as text.

    Files saved by Excel are usually UTF-8 (sometimes with an invisible "BOM"
    marker at the start) or, on older Windows setups, the Windows-1252 encoding.
    We try UTF-8 first and fall back to Windows-1252.
    """
    try:
        return csv_path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        return csv_path.read_text(encoding="cp1252")


def _read_csv_rows(path: Path) -> list[list[str]]:
    """All non-empty rows of a CSV (or plain text) file, as lists of cell texts."""
    text = _read_csv_text(path)
    # csv.Sniffer guesses the separator (",", ";" or tab) from a sample of the file.
    # European versions of Excel, for example, save CSV files with ";".
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel  # plain comma-separated as the fallback
    return [row for row in csv.reader(text.splitlines(), dialect) if any(cell.strip() for cell in row)]


def _excel_cell_text(value: object) -> str:
    """
    Turn one Excel cell into text.

    Excel sometimes turns a CAS number such as 75-05-8 into a DATE (8 May 1975)
    as soon as it is typed. The original number cannot be recovered reliably,
    so such a cell becomes "Excel date 1975-05-08". That is reported as invalid
    with an explanation (see extract_cas_numbers) and highlighted in the list,
    where it is easy to fix.
    """
    if isinstance(value, (datetime, date)):
        return f"{EXCEL_DATE_MARK} {value:%Y-%m-%d}"
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # 123.0 -> 123
    return str(value)


def _read_excel_rows(path: Path) -> list[list[str]]:
    """All non-empty rows of the first worksheet of an Excel (.xlsx) file."""
    try:
        from openpyxl import load_workbook  # only needed for Excel files
    except ImportError as error:
        raise ValueError("Reading Excel files needs the 'openpyxl' package. Install it with:\n"
                         "python -m pip install -r requirements.txt") from error

    # read_only makes big files load faster; data_only gives formula RESULTS, not formulas.
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = [[_excel_cell_text(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return [row for row in rows if any(cell.strip() for cell in row)]


def read_table_rows(path: str | Path) -> list[list[str]]:
    """Rows of a CSV, TXT or Excel file (.xlsx / .xlsm), as lists of cell texts."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return _read_excel_rows(path)
    if path.suffix.lower() == ".xls":
        raise ValueError("Old .xls files cannot be read. In Excel, use File > Save As and "
                         "choose 'Excel Workbook (*.xlsx)' or 'CSV UTF-8'.")
    return _read_csv_rows(path)


def rows_to_lines(rows: list[list[str]]) -> list[str]:
    """
    Turn table rows into lines for the list, e.g. "2530-83-8 @ Gelest".

    Which columns are used?
      * CAS: the column whose header contains "CAS" (e.g. "CAS No."). Without
        such a header, the FIRST column. If the first cell of that column has no
        digits at all (e.g. "Chemical"), it is treated as a header and skipped.
      * Manufacturer (optional): the column whose header contains "Manufacturer",
        "Supplier", "Brand", "Vendor" or "Maker".
      * Note (optional): a column whose header contains "Name" or "Chemical" is
        added after "#", so you can see which chemical each line is.
    """
    if not rows:
        return []
    header = [cell.lower() for cell in rows[0]]

    def find_column(words: tuple[str, ...]) -> int | None:
        for index, cell in enumerate(header):
            # A header cell names the column but is not itself a CAS number.
            if any(word in cell for word in words) and not _CAS_PATTERN.search(normalize_dashes(cell)):
                return index
        return None

    cas_column = find_column(("cas",))
    if cas_column is not None:
        has_header = True
    else:
        cas_column = 0
        first_cell = rows[0][0] if rows[0] else ""
        has_header = not any(character.isdigit() for character in first_cell)
    maker_column = find_column(_MANUFACTURER_HEADERS) if has_header else None
    name_column = find_column(("name", "chemical", "substance", "description")) if has_header else None
    if name_column == cas_column:
        name_column = None

    def cell(row: list[str], column: int | None) -> str:
        if column is None or column >= len(row):
            return ""
        # "@" and "#" have a special meaning in the list, and line breaks would split
        # a line, so they are removed from the manufacturer and name texts.
        return re.sub(r"[@#\r\n]+", " ", row[column]).strip()

    lines = []
    for row in rows[1:] if has_header else rows:
        cas_text = row[cas_column].strip() if cas_column < len(row) else ""
        maker, name = cell(row, maker_column), cell(row, name_column)
        if not cas_text and not maker:
            continue
        line = cas_text
        if maker:
            line += f" @ {maker}"
        if name:
            line += f"   # {name}"
        lines.append(line)
    return lines


def parse_csv_file(csv_path: str | Path) -> ParsedInput:
    """
    Read CAS numbers (and manufacturers, if there is such a column) from a CSV
    or Excel file. Comma, semicolon and tab separated files are all detected
    automatically. See rows_to_lines for which columns are used.
    """
    return parse_text_list("\n".join(rows_to_lines(read_table_rows(csv_path))))
